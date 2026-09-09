"""Template-backed onboarding document generation.

The supplied DOCX files use visual labels and generic OOXML SDTs rather than
semantic Word form controls. The manifest below records those anchors and
keeps the web fields independent from the fragile visual markup.
"""

from datetime import date, datetime, timedelta, timezone
import hashlib
import os
from pathlib import Path
import secrets
import shutil
import subprocess
from typing import Any

from docx import Document
import fitz
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import (
    OnboardingCase,
    OnboardingDocument,
    OnboardingParticipant,
    SignatureAuditLog,
    SignatureField,
    SignatureRecipient,
    SignatureRequest,
)


ORG_LEGAL_EMAIL = "legal@gobitsnbytes.org"
TEMPLATE_MANIFEST: dict[str, dict[str, Any]] = {
    "volunteer": {
        "template": "1_Volunteer_Form.docx",
        "fields": [
            {"key": "full_name", "label": "Full name", "type": "fullname", "source_anchor": "Name"},
            {"key": "date_of_birth", "label": "Date of birth", "type": "date", "source_anchor": "Date of Birth"},
            {"key": "volunteer_role", "label": "Volunteer role", "type": "text", "source_anchor": "I am a volunteer"},
            {"key": "track", "label": "Track", "type": "text", "source_anchor": "Track"},
            {"key": "consent", "label": "Volunteer consent", "type": "checkbox", "source_anchor": "I agree"},
            {"key": "signature", "label": "Signature", "type": "signature", "source_anchor": "Signature"},
            {"key": "signed_date", "label": "Signed date", "type": "date", "source_anchor": "Date"},
        ],
    },
    "parent_consent": {
        "template": "2_Parents_Consent_Fork.docx",
        "fields": [
            {"key": "minor_name", "label": "Minor name", "type": "text", "source_anchor": "child"},
            {"key": "parent_name", "label": "Parent or guardian name", "type": "fullname", "source_anchor": "Parent/Guardian"},
            {"key": "parent_email", "label": "Parent email", "type": "text", "source_anchor": "Email"},
            {"key": "consent", "label": "Parent consent", "type": "checkbox", "source_anchor": "consent"},
            {"key": "signature", "label": "Parent signature", "type": "signature", "source_anchor": "Signature"},
            {"key": "signed_date", "label": "Signed date", "type": "date", "source_anchor": "Date"},
        ],
    },
    "fork_application": {
        "template": "5_Fork_Application_Form.docx",
        "fields": [
            {"key": "fork_name", "label": "Fork name", "type": "text", "source_anchor": "Fork Name"},
            {"key": "lead_name", "label": "Fork lead", "type": "fullname", "source_anchor": "Lead"},
            {"key": "summary", "label": "Fork summary", "type": "text", "source_anchor": "Purpose"},
            {"key": "signature", "label": "Lead signature", "type": "signature", "source_anchor": "Signature"},
            {"key": "signed_date", "label": "Signed date", "type": "date", "source_anchor": "Date"},
        ],
    },
    "fork_agreement": {
        "template": "4_Fork_Agreement.docx",
        "fields": [
            {"key": "fork_name", "label": "Fork name", "type": "text", "source_anchor": "Fork"},
            {"key": "lead_name", "label": "Fork lead", "type": "fullname", "source_anchor": "Lead"},
            {"key": "agreement", "label": "Agreement accepted", "type": "checkbox", "source_anchor": "agree"},
            {"key": "signature", "label": "Lead signature", "type": "signature", "source_anchor": "Signature"},
            {"key": "signed_date", "label": "Signed date", "type": "date", "source_anchor": "Date"},
        ],
    },
    "fork_certificate": {
        "template": "3_Fork_Recognition_Certificate.docx",
        "fields": [
            {"key": "fork_name", "label": "Fork name", "type": "text", "source_anchor": "Fork"},
            {"key": "lead_name", "label": "Fork lead", "type": "text", "source_anchor": "Lead"},
            {"key": "directors", "label": "Directors", "type": "text", "source_anchor": "Director"},
            {"key": "certificate_date", "label": "Certificate date", "type": "date", "source_anchor": "Date"},
        ],
    },
    "minor_event_consent": {
        "template": "6_Minor_Consent_Event.docx",
        "fields": [
            {"key": "minor_name", "label": "Minor name", "type": "text", "source_anchor": "Minor"},
            {"key": "parent_name", "label": "Parent or guardian name", "type": "fullname", "source_anchor": "Parent"},
            {"key": "consent", "label": "Event consent", "type": "checkbox", "source_anchor": "consent"},
            {"key": "signature", "label": "Parent signature", "type": "signature", "source_anchor": "Signature"},
            {"key": "signed_date", "label": "Signed date", "type": "date", "source_anchor": "Date"},
        ],
    },
}


def age_on(date_of_birth: str, today: date | None = None) -> int:
    born = date.fromisoformat(date_of_birth)
    today = today or datetime.now(timezone.utc).date()
    return today.year - born.year - ((today.month, today.day) < (born.month, born.day))


def validate_age(date_of_birth: str, today: date | None = None) -> tuple[int, bool]:
    age = age_on(date_of_birth, today=today)
    if age < 13:
        raise ValueError("Participants must be at least 13 years old for this workflow")
    return age, age < 18


def new_portal_token() -> tuple[str, str]:
    token = secrets.token_urlsafe(32)
    return token, hashlib.sha256(token.encode()).hexdigest()


def hash_portal_token(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def template_root() -> Path:
    configured = os.getenv("ONBOARDING_TEMPLATES_DIR")
    if configured:
        return Path(configured)
    api_root = Path(__file__).resolve().parents[2]
    if (api_root / "templates").exists():
        return api_root / "templates"
    return api_root.parents[1] / "templates"


def storage_root() -> Path:
    root = Path(os.getenv("ONBOARDING_STORAGE_DIR", str(Path.cwd() / "data" / "onboarding")))
    root.mkdir(parents=True, exist_ok=True)
    return root


def manifest_for(document_key: str) -> dict[str, Any]:
    try:
        return TEMPLATE_MANIFEST[document_key]
    except KeyError as exc:
        raise ValueError(f"Unknown onboarding document type: {document_key}") from exc


def _append_evidence_section(template: Path, target: Path, values: dict[str, Any]) -> None:
    document = Document(str(template))
    document.add_page_break()
    document.add_heading("Digital completion evidence", level=1)
    document.add_paragraph("This appendix records the values entered in the onboarding portal before signature.")
    for key, value in values.items():
        document.add_paragraph(f"{key}: {value if value not in (None, '') else '[not supplied]'}")
    document.save(str(target))


def render_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    configured = os.getenv("ONBOARDING_SOFFICE_PATH")
    soffice = configured or shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice Writer is required to render onboarding DOCX files")
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError("Onboarding DOCX rendering failed")
    return pdf_path


def field_positions(field_index: int, field_type: str) -> dict[str, Any]:
    # Web fields are placed in the appended evidence page. These percentages
    # are intentionally stable; the PDF engine owns the final visual overlay.
    rows = max(0, field_index)
    return {
        "page_number": 2,
        "pos_x": 12.0,
        "pos_y": min(88.0, 12.0 + rows * 9.0),
        "width": 70.0 if field_type != "signature" else 35.0,
        "height": 6.0,
    }


async def create_signature_request(
    db: AsyncSession,
    *,
    case: OnboardingCase,
    document: OnboardingDocument,
    pdf_path: Path,
    signer_specs: list[dict[str, Any]],
    values: dict[str, Any],
    include_legal: bool = True,
) -> SignatureRequest:
    request = SignatureRequest(
        title=f"{case.title} - {document.document_key.replace('_', ' ').title()}",
        status="pending",
        original_file_path=str(pdf_path),
        created_by=case.created_by,
        expires_at=datetime.now(timezone.utc) + timedelta(days=30),
    )
    db.add(request)
    await db.flush()

    recipients: list[SignatureRecipient] = []
    for order, spec in enumerate(signer_specs, start=1):
        recipient = SignatureRecipient(
            request_id=request.id,
            name=spec["name"],
            email=spec["email"],
            role=spec.get("role", "signer"),
            signing_order=order,
            status="pending",
            access_token=secrets.token_urlsafe(24),
            requires_otp=bool(spec.get("requires_otp", False)),
            allowed_sig_type=spec.get("allowed_sig_type", "any"),
        )
        db.add(recipient)
        await db.flush()
        recipients.append(recipient)

    if include_legal:
        legal = SignatureRecipient(
            request_id=request.id,
            name="GOBITSNBYTES FOUNDATION - Legal",
            email=ORG_LEGAL_EMAIL,
            role="org_signer",
            signing_order=99,
            status="pending",
            access_token=secrets.token_urlsafe(24),
            allowed_sig_type="email_only",
        )
        db.add(legal)
        await db.flush()

    manifest = manifest_for(document.document_key)
    pdf_doc = fitz.open(str(pdf_path))
    evidence_page = max(1, len(pdf_doc))
    pdf_doc.close()
    target = recipients[0]
    for index, definition in enumerate(manifest["fields"]):
        db.add(
            SignatureField(
                request_id=request.id,
                recipient_id=target.id,
                type=definition["type"],
                page_number=evidence_page,
                pos_x=field_positions(index, definition["type"])["pos_x"],
                pos_y=field_positions(index, definition["type"])["pos_y"],
                width=field_positions(index, definition["type"])["width"],
                height=field_positions(index, definition["type"])["height"],
                required=definition["key"] in {"full_name", "signature", "consent", "parent_name", "parent_email"},
                value=str(values.get(definition["key"])) if values.get(definition["key"]) is not None else None,
            )
        )
    db.add(SignatureAuditLog(request_id=request.id, action="created", ip_address="onboarding", user_agent="motherboard", details=f"Created from onboarding case {case.id}"))
    return request


async def materialize_document(
    db: AsyncSession,
    *,
    case: OnboardingCase,
    document: OnboardingDocument,
    participant: OnboardingParticipant,
    values: dict[str, Any],
) -> SignatureRequest:
    manifest = manifest_for(document.document_key)
    template = template_root() / manifest["template"]
    if not template.exists():
        raise FileNotFoundError(f"Onboarding template is missing: {manifest['template']}")

    folder = storage_root() / str(case.id) / str(document.id)
    folder.mkdir(parents=True, exist_ok=True)
    source_path = folder / manifest["template"]
    filled_path = folder / f"filled_{manifest['template']}"
    shutil.copyfile(template, source_path)
    _append_evidence_section(source_path, filled_path, values)
    pdf_path = render_docx_to_pdf(filled_path, folder)

    document.source_docx_path = str(source_path)
    document.filled_docx_path = str(filled_path)
    document.source_pdf_path = str(pdf_path)
    document.evidence_hash = hashlib.sha256(filled_path.read_bytes()).hexdigest()
    document.field_values = values
    document.status = "awaiting_signatures"
    request = await create_signature_request(
        db,
        case=case,
        document=document,
        pdf_path=pdf_path,
        signer_specs=[{"name": participant.name, "email": participant.email}],
        values=values,
    )
    document.signature_request_id = request.id
    return request
